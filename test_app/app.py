# -*- coding: utf-8 -*-
"""Talabalar uchun onlayn test ilovasi.

Ishga tushirish:   python app.py
O'qituvchi paneli: http://<ip>:5000/admin
"""

import io
import json
import os
import random
import secrets
import sqlite3
from datetime import datetime, timedelta

from flask import (Flask, abort, g, redirect, render_template, request,
                   send_file, session, url_for)

import config
import questions as savol_moduli

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_YOL = os.path.join(BASE_DIR, "natijalar.db")

app = Flask(__name__)
app.secret_key = os.environ.get("TEST_SECRET") or secrets.token_hex(16)
app.permanent_session_lifetime = timedelta(hours=6)

SAVOLLAR = savol_moduli.yukla(os.path.join(BASE_DIR, config.SAVOLLAR_FAYLI))


# --------------------------------------------------------------------------
# Baza
# --------------------------------------------------------------------------
def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_YOL)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def db_yop(_exc):
    ulanish = g.pop("db", None)
    if ulanish is not None:
        ulanish.close()


def baza_yarat():
    ulanish = sqlite3.connect(DB_YOL)
    ulanish.executescript(
        """
        CREATE TABLE IF NOT EXISTS urinish (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            ism          TEXT NOT NULL,
            email        TEXT NOT NULL,
            guruh        TEXT,
            boshlandi    TEXT NOT NULL,
            tugadi       TEXT,
            tartib       TEXT NOT NULL,   -- JSON: [{q: id, opts: [...]}, ...]
            javoblar     TEXT,            -- JSON: {q_id: [tanlangan indekslar]}
            togri        INTEGER,
            jami         INTEGER,
            foiz         REAL,
            baho         TEXT,
            holat        TEXT NOT NULL DEFAULT 'boshlandi'
        );
        CREATE INDEX IF NOT EXISTS idx_email ON urinish(email);
        """
    )
    ulanish.commit()
    ulanish.close()


# --------------------------------------------------------------------------
# Yordamchi funksiyalar
# --------------------------------------------------------------------------
def baho_ber(foiz):
    for eng_kam, nom in sorted(config.BAHOLASH, key=lambda x: -x[0]):
        if foiz >= eng_kam:
            return nom
    return config.BAHOLASH[-1][1]


def tartib_yasa():
    """Talaba uchun savollar va variantlar tartibini tasodifiy tuzadi."""
    idlar = [s["id"] for s in SAVOLLAR]
    if config.SAVOLLARNI_ARALASHTIRISH:
        random.shuffle(idlar)
    if config.SAVOLLAR_SONI and config.SAVOLLAR_SONI < len(idlar):
        idlar = idlar[: config.SAVOLLAR_SONI]

    tartib = []
    for qid in idlar:
        opts = list(range(len(SAVOLLAR[qid]["variantlar"])))
        if config.VARIANTLARNI_ARALASHTIRISH:
            random.shuffle(opts)
        tartib.append({"q": qid, "opts": opts})
    return tartib


def urinish_ol(uid):
    qator = db().execute("SELECT * FROM urinish WHERE id = ?", (uid,)).fetchone()
    return qator


def qolgan_sekund(qator):
    boshlandi = datetime.fromisoformat(qator["boshlandi"])
    tugash = boshlandi + timedelta(minutes=config.DAVOMIYLIK_DAQIQA)
    return max(0, int((tugash - datetime.now()).total_seconds()))


def baholash(tartib, javoblar):
    """javoblar: {str(q_id): [tanlangan variantning ASL indeksi, ...]}"""
    togri = 0
    for band in tartib:
        qid = band["q"]
        kutilgan = set(SAVOLLAR[qid]["togri"])
        berilgan = set(javoblar.get(str(qid), []))
        if berilgan and berilgan == kutilgan:
            togri += 1
    return togri, len(tartib)


def yakunla(uid, javoblar):
    qator = urinish_ol(uid)
    if qator is None or qator["holat"] == "tugadi":
        return qator

    tartib = json.loads(qator["tartib"])
    togri, jami = baholash(tartib, javoblar)
    foiz = round(togri * 100.0 / jami, 1) if jami else 0.0

    db().execute(
        """UPDATE urinish SET tugadi=?, javoblar=?, togri=?, jami=?, foiz=?,
                              baho=?, holat='tugadi' WHERE id=?""",
        (datetime.now().isoformat(timespec="seconds"), json.dumps(javoblar),
         togri, jami, foiz, baho_ber(foiz), uid),
    )
    db().commit()
    return urinish_ol(uid)


# --------------------------------------------------------------------------
# Talaba sahifalari
# --------------------------------------------------------------------------
@app.route("/", methods=["GET", "POST"])
def kirish():
    xato = None
    if request.method == "POST":
        ism = (request.form.get("ism") or "").strip()
        email = (request.form.get("email") or "").strip().lower()
        guruh = (request.form.get("guruh") or "").strip()

        if len(ism) < 3:
            xato = "Ism-familiyangizni to'liq yozing."
        elif "@" not in email or "." not in email.split("@")[-1]:
            xato = "Email manzil noto'g'ri."
        elif config.RUXSAT_ETILGAN_DOMEN and not email.endswith(
            "@" + config.RUXSAT_ETILGAN_DOMEN
        ):
            xato = "Faqat @{} pochtasi qabul qilinadi.".format(config.RUXSAT_ETILGAN_DOMEN)

        if xato is None:
            mavjud = db().execute(
                "SELECT * FROM urinish WHERE email = ? ORDER BY id DESC LIMIT 1", (email,)
            ).fetchone()

            if mavjud is not None:
                if mavjud["holat"] == "tugadi":
                    if config.BIR_MARTA_TOPSHIRISH:
                        return render_template(
                            "xabar.html", sarlavha="Test allaqachon topshirilgan",
                            matn="Bu email bilan test bir marta topshirilgan. "
                                 "Qayta topshirish mumkin emas.")
                elif qolgan_sekund(mavjud) > 0:
                    # Tugallanmagan urinish -- davom ettiramiz
                    session["uid"] = mavjud["id"]
                    return redirect(url_for("test"))
                else:
                    yakunla(mavjud["id"], json.loads(mavjud["javoblar"] or "{}"))
                    if config.BIR_MARTA_TOPSHIRISH:
                        return render_template(
                            "xabar.html", sarlavha="Vaqt tugagan",
                            matn="Oldingi urinishingizda vaqt tugab qolgan va test "
                                 "avtomatik yakunlangan.")

            tartib = tartib_yasa()
            kursor = db().execute(
                "INSERT INTO urinish (ism, email, guruh, boshlandi, tartib, javoblar)"
                " VALUES (?,?,?,?,?,?)",
                (ism, email, guruh, datetime.now().isoformat(timespec="seconds"),
                 json.dumps(tartib), "{}"),
            )
            db().commit()
            session["uid"] = kursor.lastrowid
            return redirect(url_for("test"))

    return render_template(
        "kirish.html", xato=xato, test_nomi=config.TEST_NOMI,
        daqiqa=config.DAVOMIYLIK_DAQIQA,
        savol_soni=config.SAVOLLAR_SONI or len(SAVOLLAR),
        domen=config.RUXSAT_ETILGAN_DOMEN,
    )


@app.route("/test")
def test():
    uid = session.get("uid")
    qator = urinish_ol(uid) if uid else None
    if qator is None:
        return redirect(url_for("kirish"))
    if qator["holat"] == "tugadi":
        return redirect(url_for("natija"))

    qolgan = qolgan_sekund(qator)
    if qolgan <= 0:
        yakunla(uid, json.loads(qator["javoblar"] or "{}"))
        return redirect(url_for("natija"))

    tartib = json.loads(qator["tartib"])
    korinish = []
    for n, band in enumerate(tartib, start=1):
        s = SAVOLLAR[band["q"]]
        korinish.append({
            "n": n,
            "qid": band["q"],
            "matn": s["matn"],
            "kop": len(s["togri"]) > 1,
            "variantlar": [{"idx": i, "matn": s["variantlar"][i]} for i in band["opts"]],
        })

    return render_template("test.html", test_nomi=config.TEST_NOMI,
                           savollar=korinish, qolgan=qolgan, ism=qator["ism"])


@app.route("/topshirish", methods=["POST"])
def topshirish():
    uid = session.get("uid")
    if not uid:
        return redirect(url_for("kirish"))

    javoblar = {}
    for kalit in request.form:
        if kalit.startswith("q_"):
            qid = kalit[2:]
            javoblar[qid] = [int(v) for v in request.form.getlist(kalit)]

    yakunla(uid, javoblar)
    return redirect(url_for("natija"))


@app.route("/natija")
def natija():
    uid = session.get("uid")
    qator = urinish_ol(uid) if uid else None
    if qator is None:
        return redirect(url_for("kirish"))
    if qator["holat"] != "tugadi":
        return redirect(url_for("test"))

    session.pop("uid", None)

    if not config.NATIJA_DARHOL_KORINSIN:
        return render_template("xabar.html", sarlavha="Test yakunlandi",
                               matn="Javoblaringiz qabul qilindi. Natija o'qituvchi "
                                    "tomonidan e'lon qilinadi.")

    xatolar = []
    if config.XATOLAR_KORSATILSIN:
        javoblar = json.loads(qator["javoblar"] or "{}")
        for band in json.loads(qator["tartib"]):
            qid = band["q"]
            s = SAVOLLAR[qid]
            berilgan = set(javoblar.get(str(qid), []))
            if berilgan != set(s["togri"]):
                xatolar.append({
                    "matn": s["matn"],
                    "sizniki": [s["variantlar"][i] for i in sorted(berilgan)] or ["(javob yo'q)"],
                    "togri": [s["variantlar"][i] for i in sorted(s["togri"])],
                })

    return render_template("natija.html", u=qator, xatolar=xatolar,
                           test_nomi=config.TEST_NOMI)


# --------------------------------------------------------------------------
# O'qituvchi paneli
# --------------------------------------------------------------------------
def admin_tekshir():
    if session.get("admin"):
        return True
    return False


@app.route("/admin", methods=["GET", "POST"])
def admin():
    if request.method == "POST" and not admin_tekshir():
        if request.form.get("parol") == config.ADMIN_PAROL:
            session["admin"] = True
            return redirect(url_for("admin"))
        return render_template("admin_kirish.html", xato="Parol noto'g'ri.")

    if not admin_tekshir():
        return render_template("admin_kirish.html", xato=None)

    qatorlar = db().execute(
        "SELECT * FROM urinish ORDER BY id DESC"
    ).fetchall()
    tugagan = [q for q in qatorlar if q["holat"] == "tugadi"]
    ortacha = round(sum(q["foiz"] for q in tugagan) / len(tugagan), 1) if tugagan else 0
    return render_template("admin.html", qatorlar=qatorlar, ortacha=ortacha,
                           test_nomi=config.TEST_NOMI, savol_soni=len(SAVOLLAR))


@app.route("/admin/chiqish")
def admin_chiqish():
    session.pop("admin", None)
    return redirect(url_for("admin"))


@app.route("/admin/ochirish/<int:uid>", methods=["POST"])
def admin_ochirish(uid):
    if not admin_tekshir():
        abort(403)
    db().execute("DELETE FROM urinish WHERE id = ?", (uid,))
    db().commit()
    return redirect(url_for("admin"))


def _natijalar():
    return db().execute(
        "SELECT * FROM urinish WHERE holat='tugadi' ORDER BY foiz DESC, ism"
    ).fetchall()


def _davomiylik(q):
    if not q["tugadi"]:
        return ""
    fark = datetime.fromisoformat(q["tugadi"]) - datetime.fromisoformat(q["boshlandi"])
    d, s = divmod(int(fark.total_seconds()), 60)
    return "{} daq {} sek".format(d, s)


@app.route("/admin/excel")
def admin_excel():
    if not admin_tekshir():
        abort(403)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"

    sarlavhalar = ["№", "Ism-familiya", "Email", "Guruh", "Boshlandi", "Tugadi",
                   "Sarflangan vaqt", "To'g'ri", "Jami", "Foiz %", "Baho"]
    ws.append(sarlavhalar)
    for hujayra in ws[1]:
        hujayra.font = Font(bold=True, color="FFFFFF")
        hujayra.fill = PatternFill("solid", fgColor="4472C4")
        hujayra.alignment = Alignment(horizontal="center", vertical="center")

    for n, q in enumerate(_natijalar(), start=1):
        ws.append([n, q["ism"], q["email"], q["guruh"] or "", q["boshlandi"],
                   q["tugadi"], _davomiylik(q), q["togri"], q["jami"],
                   q["foiz"], q["baho"]])

    for kenglik, ustun in zip([5, 28, 30, 10, 20, 20, 16, 9, 7, 9, 16], "ABCDEFGHIJK"):
        ws.column_dimensions[ustun].width = kenglik
    ws.freeze_panes = "A2"

    # Ikkinchi varaq: har bir savol bo'yicha javoblar
    ws2 = wb.create_sheet("Javoblar")
    ws2.append(["Ism-familiya", "Email", "Savol", "Talaba javobi", "To'g'ri javob", "Natija"])
    for hujayra in ws2[1]:
        hujayra.font = Font(bold=True)
    for q in _natijalar():
        javoblar = json.loads(q["javoblar"] or "{}")
        for band in json.loads(q["tartib"]):
            s = SAVOLLAR[band["q"]]
            berilgan = sorted(javoblar.get(str(band["q"]), []))
            ws2.append([
                q["ism"], q["email"], s["matn"],
                "; ".join(s["variantlar"][i] for i in berilgan) or "(javob yo'q)",
                "; ".join(s["variantlar"][i] for i in sorted(s["togri"])),
                "to'g'ri" if set(berilgan) == set(s["togri"]) else "xato",
            ])
    for kenglik, ustun in zip([26, 28, 50, 32, 32, 10], "ABCDEF"):
        ws2.column_dimensions[ustun].width = kenglik

    # Uchinchi varaq: savollar statistikasi
    ws3 = wb.create_sheet("Savollar tahlili")
    ws3.append(["Savol", "To'g'ri javob berganlar", "Jami urinish", "To'g'ri %"])
    for hujayra in ws3[1]:
        hujayra.font = Font(bold=True)
    hisob = {s["id"]: [0, 0] for s in SAVOLLAR}
    for q in _natijalar():
        javoblar = json.loads(q["javoblar"] or "{}")
        for band in json.loads(q["tartib"]):
            qid = band["q"]
            hisob[qid][1] += 1
            if set(javoblar.get(str(qid), [])) == set(SAVOLLAR[qid]["togri"]):
                hisob[qid][0] += 1
    for s in SAVOLLAR:
        yaxshi, jami = hisob[s["id"]]
        ws3.append([s["matn"], yaxshi, jami,
                    round(yaxshi * 100.0 / jami, 1) if jami else 0])
    for kenglik, ustun in zip([60, 22, 14, 12], "ABCD"):
        ws3.column_dimensions[ustun].width = kenglik

    bufer = io.BytesIO()
    wb.save(bufer)
    bufer.seek(0)
    return send_file(bufer, as_attachment=True,
                     download_name="test_natijalari_{}.xlsx".format(
                         datetime.now().strftime("%Y-%m-%d_%H%M")),
                     mimetype="application/vnd.openxmlformats-officedocument."
                              "spreadsheetml.sheet")


@app.route("/admin/word")
def admin_word():
    if not admin_tekshir():
        abort(403)
    from docx import Document
    from docx.shared import Pt

    hujjat = Document()
    hujjat.add_heading(config.TEST_NOMI, level=0)
    hujjat.add_paragraph("Hisobot sanasi: " + datetime.now().strftime("%d.%m.%Y %H:%M"))

    natijalar = _natijalar()
    hujjat.add_heading("Umumiy natijalar", level=1)
    jadval = hujjat.add_table(rows=1, cols=7)
    jadval.style = "Light Grid Accent 1"
    for hujayra, matn in zip(jadval.rows[0].cells,
                             ["№", "Ism-familiya", "Email", "Guruh", "To'g'ri", "Foiz", "Baho"]):
        hujayra.text = matn
        hujayra.paragraphs[0].runs[0].font.bold = True

    for n, q in enumerate(natijalar, start=1):
        hujayralar = jadval.add_row().cells
        qiymatlar = [str(n), q["ism"], q["email"], q["guruh"] or "-",
                     "{}/{}".format(q["togri"], q["jami"]),
                     "{}%".format(q["foiz"]), q["baho"]]
        for hujayra, matn in zip(hujayralar, qiymatlar):
            hujayra.text = matn
            hujayra.paragraphs[0].runs[0].font.size = Pt(10)

    hujjat.add_page_break()
    hujjat.add_heading("Har bir talabaning javoblari", level=1)
    for q in natijalar:
        hujjat.add_heading("{} ({}) - {}, {}".format(
            q["ism"], q["email"], q["baho"], "{}%".format(q["foiz"])), level=2)
        javoblar = json.loads(q["javoblar"] or "{}")
        for n, band in enumerate(json.loads(q["tartib"]), start=1):
            s = SAVOLLAR[band["q"]]
            berilgan = sorted(javoblar.get(str(band["q"]), []))
            ok = set(berilgan) == set(s["togri"])
            p = hujjat.add_paragraph("{}. {}".format(n, s["matn"]))
            p.runs[0].font.bold = True
            hujjat.add_paragraph("Javobi: {}   [{}]".format(
                "; ".join(s["variantlar"][i] for i in berilgan) or "(javob yo'q)",
                "to'g'ri" if ok else "xato"), style="List Bullet")
            if not ok:
                hujjat.add_paragraph("To'g'ri javob: " + "; ".join(
                    s["variantlar"][i] for i in sorted(s["togri"])), style="List Bullet")

    bufer = io.BytesIO()
    hujjat.save(bufer)
    bufer.seek(0)
    return send_file(bufer, as_attachment=True,
                     download_name="test_natijalari_{}.docx".format(
                         datetime.now().strftime("%Y-%m-%d_%H%M")),
                     mimetype="application/vnd.openxmlformats-officedocument."
                              "wordprocessingml.document")


# --------------------------------------------------------------------------
if __name__ == "__main__":
    baza_yarat()
    import socket

    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
    except OSError:
        ip = "127.0.0.1"

    print("\n" + "=" * 60)
    print("  {}  ({} ta savol)".format(config.TEST_NOMI, len(SAVOLLAR)))
    print("=" * 60)
    print("  Talabalar uchun havola :  http://{}:5000".format(ip))
    print("  O'qituvchi paneli      :  http://{}:5000/admin".format(ip))
    print("  Admin parol            :  {}".format(config.ADMIN_PAROL))
    print("=" * 60 + "\n")

    app.run(host="0.0.0.0", port=5000, debug=False)
